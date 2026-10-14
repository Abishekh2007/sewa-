# Fills the official Seva First Innovation Challenge 2026 template for RAKSHA-PATH X.
# Keeps template logos, titles, team-name oval and footer; replaces the pointer
# text boxes with structured content that answers each pointer verbatim.
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = r"C:\Users\thena\Desktop\sewa\raksha-path-x"
SRC = ROOT + r"\Seva First Innovation Challenge 2026.pptx"
OUT = ROOT + r"\presentation\RAKSHA-PATH X - Seva First Innovation Challenge 2026.pptx"
SHOT = ROOT + r"\docs\assets\prototype-dashboard.jpg"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x55, 0x97)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x59, 0x59, 0x59)
ORANGE = RGBColor(0xE4, 0x6C, 0x0A)
GREEN = RGBColor(0x00, 0x8A, 0x3E)
TINT = RGBColor(0xEE, 0xF2, 0xF8)
TINT_O = RGBColor(0xFD, 0xF0, 0xE6)
LINE = RGBColor(0xC9, 0xD3, 0xE3)
PLACE = RGBColor(0xC0, 0x00, 0x00)  # red = placeholder the team must fill
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
STATUS = [  # (label, rgb, rule)
    ("Verified passable", RGBColor(0x2E, 0x8B, 0x57), "Recent evidence from an authenticated or verified source"),
    ("Potential hazard", RGBColor(0xD9, 0x9A, 0x1E), "Hazard reported, not yet verified"),
    ("Blocked", RGBColor(0xC0, 0x39, 0x2B), "Confirmed; never used for routing"),
    ("Unknown", RGBColor(0x80, 0x80, 0x80), "No adequate evidence (no report is not 'open')"),
    ("Stale evidence", RGBColor(0x6A, 0x6F, 0xB0), "Evidence older than its hazard-specific limit"),
    ("Conflicting evidence", RGBColor(0xD8, 0x69, 0x2A), "Reports disagree; both kept, verifier tasked"),
]

prs = Presentation(SRC)
S = prs.slides


def para_fmt(p, size, bold=False, color=INK, italic=False, align=None, space_after=4):
    if align is not None:
        p.alignment = align
    p.space_after = Pt(space_after)
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = "Arial"
        r.font.color.rgb = color


def set_bullet(p, indent_in=0.18):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bc = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bf)
    pPr.append(bc)


def textbox(slide, x, y, w, h, items, anchor=MSO_ANCHOR.TOP, margin=0.04):
    """items: list of dicts {runs:[(text, {bold,color,italic})] | text, size, bold, color, bullet, align, after}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        runs = it.get("runs") or [(it["text"], {})]
        for text, ro in runs:
            r = p.add_run()
            r.text = text
            r.font.name = "Arial"
            r.font.size = Pt(ro.get("size", it.get("size", 14)))
            r.font.bold = ro.get("bold", it.get("bold", False))
            r.font.italic = ro.get("italic", it.get("italic", False))
            r.font.color.rgb = ro.get("color", it.get("color", INK))
        if it.get("align") is not None:
            p.alignment = it["align"]
        p.space_after = Pt(it.get("after", 4))
        if it.get("bullet"):
            set_bullet(p)
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def heading(slide, x, y, w, text, size=16, color=BLUE):
    return textbox(slide, x, y, w, 0.4, [{"runs": [("❖ ", {"color": color}), (text, {})], "size": size, "bold": True, "color": color}])


def shape_named(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_shape_text(shape, text, color=None):
    """Replace text keeping the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    runs = [r for p in tf.paragraphs for r in p.runs]
    keep = runs[0] if runs else p0.add_run()
    keep.text = text
    if color is not None:
        keep.font.color.rgb = color
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)


# ---------- common: team-name oval + footer on content slides ----------
for slide in list(S)[1:6]:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
            set_shape_text(sh, "[Team Name]", PLACE)
        elif sh.has_text_frame and "Template" in sh.text_frame.text:
            set_shape_text(sh, "SEVA FIRST INNOVATION CHALLENGE 2026")

# ---------- slide 1: title page ----------
s1 = S[0]
tb = shape_named(s1, "TextBox 9")
tb.width = Inches(7.4)
values = {
    "Problem Statement ID": ("[PS ID]", True),
    "Problem Statement Title": ("[PS Title]", True),
    "Theme": ("Surakshit Bharat (Theme 5)", False),
    "PS Category": None,
    "Team ID": ("[Team ID]", True),
    "Team Name": ("[Team Name]", True),
}
for p in tb.text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(20)
    txt = "".join(r.text for r in p.runs)
    for key, val in values.items():
        if txt.strip().startswith(key):
            if key == "PS Category":
                # keep label, set category
                p.runs[0].text = "PS Category- "
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                r = p.add_run(); r.text = "Software"
                r.font.bold = False; r.font.size = p.runs[0].font.size; r.font.name = "Arial"
                r.font.color.rgb = INK
            else:
                text, is_place = val
                label = txt.rstrip()
                if not label.endswith(("–", "-")):
                    label = label + " -"
                p.runs[0].text = label + " "
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                r = p.add_run(); r.text = text
                r.font.bold = False; r.font.size = Pt(16); r.font.name = "Arial"
                r.font.color.rgb = PLACE if is_place else INK
            break

# ---------- slide 2: idea title / proposed solution ----------
s2 = S[1]
set_shape_text(shape_named(s2, "Title 1"), "RAKSHA-PATH X")
textbox(s2, 2.3, 1.22, 8.8, 0.45, [{"text": "Proof-of-Passability Emergency Intelligence Network  ·  “Every route needs evidence. Every rescue needs a safer path.”", "size": 14, "italic": True, "color": NAVY, "align": PP_ALIGN.CENTER}])
body = shape_named(s2, "TextBox 8")
# keep only the template's section header paragraph
for p in list(body.text_frame.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
for r in body.text_frame.paragraphs[0].runs:
    r.font.size = Pt(26)
body.top = Inches(1.95); body.height = Inches(0.55); body.left = Inches(0.2); body.width = Inches(13.0)

L = 0.45; W = 7.55
y = 2.62
blocks = [
    ("Detailed explanation of the proposed solution", [
        "A decision-support platform for emergency control rooms that rates every road segment by the evidence behind it: recency, source trust, verification",
        "Evidence-aware routing for ambulances, fire tenders and rescue units, with alternatives, vehicle restrictions and field-verification tasks",
    ]),
    ("How it addresses the problem", [
        "Navigation apps assume roads are open. In floods and landslides they are not, and the cost is lost rescue time",
        "Never routes through a confirmed block; when no adequately verified route exists it says so and proposes the nearest reachable meeting point",
    ]),
    ("Innovation and uniqueness of the solution", [
        "Proof-of-Passability: evidence expires per hazard type (flood 2 h, fallen tree 12 h); “no report” is treated as Unknown, never “open”",
        "Transparent “Why this status?” trail instead of a black-box AI score; offline-first; humans make the final call",
    ]),
]
for head, bullets in blocks:
    heading(s2, L - 0.05, y, W, head, size=15)
    y += 0.36
    textbox(s2, L + 0.1, y, W - 0.1, 0.95, [{"text": b, "size": 12.5, "bullet": True, "after": 3} for b in bullets])
    y += 0.98

# right panel: six evidence-based road statuses
px, py, pw = 8.35, 2.62, 4.6
box(s2, px, py, pw, 4.2, fill=TINT)
textbox(s2, px + 0.2, py + 0.12, pw - 0.4, 0.35, [{"text": "Every road gets an evidence-based status", "size": 13, "bold": True, "color": NAVY}])
ry = py + 0.55
for label, rgb, rule in STATUS:
    chip = box(s2, px + 0.2, ry + 0.06, 0.22, 0.22, fill=rgb, shape=MSO_SHAPE.OVAL)
    textbox(s2, px + 0.55, ry - 0.02, pw - 0.75, 0.55, [
        {"text": label, "size": 12, "bold": True, "color": INK, "after": 0},
        {"text": rule, "size": 10.5, "color": MUTED, "after": 0},
    ])
    ry += 0.6

# ---------- slide 3: technical approach ----------
s3 = S[2]
remove(shape_named(s3, "TextBox 8"))
heading(s3, 0.4, 2.12, 8.4, "Technologies to be used (e.g. programming languages, frameworks, hardware)", size=14)
tech = [
    ("Frontend", "React 18 + TypeScript, Vite, Tailwind CSS (responsive, keyboard-accessible)"),
    ("Map", "Custom SVG road-network renderer; no paid map or API keys"),
    ("Routing", "Dijkstra graph engine in TypeScript: 3 modes, alternatives, last-mile"),
    ("Evidence", "Rule engine: freshness, trust tiers, duplicates, conflicts"),
    ("Data", "LocalStorage persistence, Service Worker offline cache, CSV export"),
    ("Quality", "Vitest: 34 automated tests, 27 evaluation cases, 8 demo scenarios"),
    ("Hardware (opt.)", "ESP32 + water-level sensor sending {sensorId, valueCm, ts} over Wi-Fi"),
]
ty = 2.55
for k, v in tech:
    textbox(s3, 0.55, ty, 8.1, 0.32, [{"runs": [(k + ":  ", {"bold": True, "color": NAVY}), (v, {})], "size": 12}])
    ty += 0.33
pic = s3.shapes.add_picture(SHOT, Inches(9.05), Inches(1.2), height=Inches(3.5))
pic.left = Inches(12.95) - pic.width
pic.line.color.rgb = LINE
textbox(s3, 8.75, 4.72, 4.2, 0.3, [{"text": "Working prototype: operations dashboard (simulated data)", "size": 10, "italic": True, "color": MUTED, "align": PP_ALIGN.RIGHT}])

heading(s3, 0.4, 5.08, 12.5, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", size=14)
flow = [
    ("Inputs", "Crowd reports, field inspections, sensor readings"),
    ("Evidence engine", "Freshness · trust tier · duplicates · conflicts"),
    ("Road status", "6 statuses + “Why this status?” explanation"),
    ("Routing", "Fastest · Evidence-prioritized · Conservative"),
    ("Decision output", "Route + alternatives, or “no verified route” + meeting point"),
    ("Field verification", "Verifier confirms; result updates graph and reroutes"),
]
fx, fw, gap, fy, fh = 0.35, 1.86, 0.28, 5.52, 1.28
for i, (t, d) in enumerate(flow):
    x = fx + i * (fw + gap)
    box(s3, x, fy, fw, fh, fill=TINT if i % 5 else TINT_O, line=LINE)
    textbox(s3, x + 0.06, fy + 0.06, fw - 0.12, fh - 0.12, [
        {"text": t, "size": 12, "bold": True, "color": NAVY, "after": 2},
        {"text": d, "size": 10.5, "color": INK, "after": 0},
    ])
    if i < len(flow) - 1:
        a = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + fw + 0.04), Inches(fy + fh / 2 - 0.1), Inches(gap - 0.08), Inches(0.2))
        a.fill.solid(); a.fill.fore_color.rgb = ORANGE; a.line.fill.background(); a.shadow.inherit = False

# ---------- slide 4: feasibility ----------
s4 = S[3]
remove(shape_named(s4, "TextBox 8"))
heading(s4, 0.4, 2.12, 4.9, "Analysis of the feasibility of the idea", size=14)
stats = [("Working", "prototype built and demonstrable today"), ("₹0", "software licences, API keys or paid maps"), ("~₹2,600", "per optional sensor node (estimate*)"), ("27/27", "evaluation cases pass in the test suite")]
sy = 2.55
for big, small in stats:
    box(s4, 0.5, sy, 4.7, 0.78, fill=TINT)
    textbox(s4, 0.62, sy + 0.05, 1.75, 0.68, [{"text": big, "size": 20 if len(big) < 7 else 17, "bold": True, "color": NAVY}], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s4, 2.35, sy + 0.05, 2.8, 0.68, [{"text": small, "size": 11.5, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)
    sy += 0.88
textbox(s4, 0.5, 6.1, 4.7, 0.7, [
    {"text": "Runs offline on an ordinary laptop; open-source stack; one-click Windows launcher.", "size": 11, "color": INK, "after": 2},
    {"text": "*Indicative component estimates; require supplier verification.", "size": 9.5, "italic": True, "color": MUTED},
])

rows = [
    ("False or prank crowd reports", "Trust tiers; anonymous reports stay unverified; duplicate detection; verification tasks"),
    ("Outdated information", "Hazard-specific expiry; stale roads flagged and queued for re-inspection"),
    ("Network loss during disasters", "Offline-first: reports queued, synced later with conflict detection"),
    ("Over-reliance on automation", "Decision support only; human confirmation; never labels a route “safe”"),
    ("Sensor faults or misreadings", "A sensor alone never marks a road passable; out-of-range and disconnect alerts"),
    ("Real road data & approvals", "Pilot with local authority GIS / OpenStreetMap data under formal permission"),
]
tx, tyy, tw = 5.55, 2.12, 7.4
tbl = s4.shapes.add_table(len(rows) + 1, 2, Inches(tx), Inches(tyy), Inches(tw), Inches(4.6)).table
tbl.columns[0].width = Inches(2.75)
tbl.columns[1].width = Inches(tw - 2.75)
hdr = ["Potential challenges and risks", "Strategies for overcoming these challenges"]
tbl.rows[0].height = Inches(0.5)
for i in range(1, len(rows) + 1):
    tbl.rows[i].height = Inches(0.68)
for c, h in enumerate(hdr):
    cell = tbl.cell(0, c)
    cell.text = h
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    para_fmt(cell.text_frame.paragraphs[0], 12.5, bold=True, color=WHITE, space_after=0)
for r, (a, b) in enumerate(rows, start=1):
    for c, t in enumerate((a, b)):
        cell = tbl.cell(r, c)
        cell.text = t
        cell.fill.solid(); cell.fill.fore_color.rgb = TINT if r % 2 else WHITE
        para_fmt(cell.text_frame.paragraphs[0], 11, bold=(c == 0), color=NAVY if c == 0 else INK, space_after=0)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

# ---------- slide 5: impact ----------
s5 = S[4]
remove(shape_named(s5, "TextBox 8"))
heading(s5, 0.4, 2.12, 6.2, "Potential impact on the target audience", size=14)
aud = [
    "District Emergency Operation Centres and control rooms",
    "Ambulance, fire and rescue crews (SDRF / NDRF teams)",
    "Local authorities and ward offices verifying roads",
    "Residents of flood- and landslide-prone areas",
]
textbox(s5, 0.55, 2.52, 6.0, 1.6, [{"text": t, "size": 12.5, "bullet": True, "after": 3} for t in aud])
# measured prototype results
box(s5, 6.85, 2.12, 6.1, 2.05, fill=TINT)
textbox(s5, 7.05, 2.2, 5.7, 0.32, [{"text": "Measured in the prototype vs. a plain shortest-path baseline", "size": 11.5, "bold": True, "color": NAVY}])
res = [("4 → 0", "confirmed-blocked road segments routed through"), ("7 → 0", "unverified segments on chosen routes"), ("8 / 8", "end-to-end demo scenarios pass")]
for i, (big, small) in enumerate(res):
    x = 7.05 + i * 1.95
    textbox(s5, x, 2.55, 1.85, 0.55, [{"text": big, "size": 24, "bold": True, "color": ORANGE if i < 2 else GREEN}])
    textbox(s5, x, 3.1, 1.8, 0.7, [{"text": small, "size": 10.5, "color": INK}])
textbox(s5, 7.05, 3.82, 5.8, 0.3, [{"text": "11 routing test cases on simulated data; not field validation.", "size": 9.5, "italic": True, "color": MUTED}])

heading(s5, 0.4, 4.35, 12.5, "Benefits of the solution (social, economic, environmental, etc.)", size=14)
cards = [
    ("Social", "Fewer rescue vehicles sent into blocked or flooded roads; faster verified access to people in need"),
    ("Economic", "Zero licence cost; runs on existing laptops and phones; low-cost optional sensors"),
    ("Environmental", "Fewer wasted detours and idle trips; reusable flood-level data for planning"),
    ("Governance", "Every decision explainable and logged; supports accountable, auditable response"),
]
cw, cg, cy = 3.0, 0.17, 4.78
for i, (t, d) in enumerate(cards):
    x = 0.45 + i * (cw + cg)
    box(s5, x, cy, cw, 1.85, fill=TINT_O if i == 0 else TINT)
    textbox(s5, x + 0.15, cy + 0.12, cw - 0.3, 1.75, [
        {"text": t, "size": 13.5, "bold": True, "color": NAVY, "after": 4},
        {"text": d, "size": 13, "color": INK},
    ])

# ---------- slide 6: research & references ----------
s6 = S[5]
remove(shape_named(s6, "TextBox 8"))
heading(s6, 0.4, 2.12, 12.5, "Details / Links of the reference and research work", size=14)
box(s6, 0.45, 2.6, 5.3, 3.55, fill=TINT)
textbox(s6, 0.65, 2.7, 4.95, 0.35, [{"text": "Research basis for the design", "size": 13, "bold": True, "color": NAVY}])
insights = [
    "Consumer navigation optimises travel time from traffic data; it does not certify that a road is physically passable during floods or landslides",
    "Disaster reports are often crowd-sourced, unverified and quickly outdated, so evidence needs freshness and trust levels",
    "Emergency decisions must be explainable and auditable, so we use transparent rules, not black-box scores",
    "Connectivity degrades in disasters, so the system has to work offline first",
]
textbox(s6, 0.65, 3.1, 4.95, 3.0, [{"text": t, "size": 12, "bullet": True, "after": 5} for t in insights])

refs = [
    ("NDMA, National Disaster Management Plan (2019)", "https://ndma.gov.in"),
    ("NDMA, Guidelines on Management of Urban Flooding (2010)", "https://ndma.gov.in"),
    ("UNDRR, Sendai Framework for Disaster Risk Reduction 2015–2030", "https://www.undrr.org"),
    ("Central Water Commission, flood forecasting services", "https://cwc.gov.in"),
    ("Dijkstra, E. W. (1959). A note on two problems in connexion with graphs. Numerische Mathematik 1, 269–271", None),
    ("Humanitarian OpenStreetMap Team, open map data for disaster response", "https://www.hotosm.org"),
]
items = [{"text": "References", "size": 13, "bold": True, "color": NAVY, "after": 6}]
for i, (t, link) in enumerate(refs, 1):
    runs = [(f"{i}. ", {"bold": True, "color": NAVY}), (t, {})]
    if link:
        runs.append(("  " + link, {"color": BLUE}))
    items.append({"runs": runs, "size": 11, "after": 5})
items.append({"runs": [("7. ", {"bold": True, "color": NAVY}), ("[Add any other papers / articles the team consulted]", {"color": PLACE})], "size": 11, "after": 10})
items.append({"text": "Prototype", "size": 13, "bold": True, "color": NAVY, "after": 4})
items.append({"runs": [("Source code: ", {"bold": True}), ("[GitHub repository URL]", {"color": PLACE})], "size": 11, "after": 3})
items.append({"runs": [("Demo video: ", {"bold": True}), ("[Demo video URL]", {"color": PLACE})], "size": 11, "after": 3})
textbox(s6, 6.05, 2.6, 6.9, 4.2, items)

# ---------- drop the template's "Important instructions" slide ----------
sldIdLst = prs.slides._sldIdLst
last = list(sldIdLst)[6]
prs.part.drop_rel(last.get(qn("r:id")))
sldIdLst.remove(last)

prs.save(OUT)
print("saved", OUT)
